#!/usr/bin/env python3
"""
Generate the LOMAMO Foundation master patient workbook template.

Produces a styled, multi-sheet XLSX that the offline import engine
(tools/import_engine.py) understands, with a field dictionary and
three synthetic demonstration rows that must be deleted before use.

    python3 tools/make_master_template.py
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "examples"
OUT.mkdir(exist_ok=True)
FILE = OUT / "LOMAMO-Master-Patient-Workbook-TEMPLATE.xlsx"

RED = "C8102E"
INK = "15171A"
CREAM = "FBF7F2"
GREY = "F4F6F9"

thin = Side(style="thin", color="D5DBE3")
border = Border(left=thin, right=thin, top=thin, bottom=thin)


def title_block(ws, title, subtitle, width):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    c = ws.cell(row=1, column=1, value="LOMAMO FOUNDATION KENYA")
    c.font = Font(name="Calibri", size=15, bold=True, color=RED)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    c = ws.cell(row=2, column=1, value=title)
    c.font = Font(name="Calibri", size=12, bold=True, color=INK)

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=width)
    c = ws.cell(row=3, column=1, value=subtitle)
    c.font = Font(name="Calibri", size=9, italic=True, color="5F6B7A")


def header_row(ws, row, headers, widths=None):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(name="Calibri", size=9, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=INK)
        c.alignment = Alignment(vertical="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[row].height = 30
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def style_data(ws, first_row, last_row, cols):
    for r in range(first_row, last_row + 1):
        for c in range(1, cols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
            cell.font = Font(name="Calibri", size=9)
            cell.alignment = Alignment(vertical="top", wrap_text=(c > 6))


def build():
    wb = Workbook()

    # ---------------- Instructions ----------------
    ws = wb.active
    ws.title = "Instructions"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110

    ws["B2"] = "LOMAMO FOUNDATION KENYA"
    ws["B2"].font = Font(size=18, bold=True, color=RED)
    ws["B3"] = "Master Patient Workbook — template"
    ws["B3"].font = Font(size=13, bold=True, color=INK)
    ws["B4"] = "Love Matters Most · Serving our neighbours in love"
    ws["B4"].font = Font(size=9, italic=True, color="5F6B7A")

    notes = [
        ("Purpose", "This workbook is the organisation's working record of patients, home visits and hospital referrals. "
                    "It is the source for the first major data operation in CareLink DIMS."),
        ("How to use", "Complete one row per patient on the 'Patient Register' sheet. Never merge cells, never insert "
                       "blank rows between records, and do not change the header row. Additional columns are permitted — "
                       "they are preserved in the original file and reported as unmapped."),
        ("Demonstration rows", "The three rows marked DEMO are synthetic samples showing the expected format. "
                               "Delete them before importing live data. They are not real patients."),
        ("Offline processing", "Run:  python3 tools/import_engine.py --workbook \"this file.xlsx\" --out lomamo_master.json "
                              "--report reconciliation.md\nThe workbook is read on this device only. Nothing is uploaded."),
        ("Confidentiality", "This file contains patient health information. Store it on an encrypted volume, restrict access "
                            "to authorised staff, and never send it by ordinary email or public messaging services."),
        ("Duplicate safety", "Identical names are never merged automatically. Supply a File Number and CCC Number wherever "
                             "they exist so that records can be matched reliably."),
        ("Dates", "Use YYYY-MM-DD (for example 2026-03-14). Excel date cells are also accepted and are converted on import."),
        ("Programme status", "Use one of: Active, Enrolled, Awaiting verification, Discharged. A discharge date overrides a "
                             "conflicting status and is flagged in the reconciliation report."),
        ("Ministry areas", "Kisumu: Nyalenda, Nairobi Road, Manyatta, Ahero, Ombeyi, Chiga, Nyahera. "
                           "Homa Bay: Makongeni, Sofia, Shauri Yako (list pending verification — add further areas in "
                           "DIMS Organisation Configuration before importing them)."),
        ("Support", "Contact the Data Management & M&E officer or the Chief System Administrator before altering the "
                    "structure of this workbook."),
    ]
    r = 6
    for head, body in notes:
        c = ws.cell(row=r, column=2, value=head)
        c.font = Font(size=10, bold=True, color=RED)
        r += 1
        c = ws.cell(row=r, column=2, value=body)
        c.font = Font(size=10, color="2B3038")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[r].height = 30 if len(body) < 160 else 44
        r += 2

    # ---------------- Patient Register ----------------
    ws = wb.create_sheet("Patient Register")
    headers = [
        "Patient Name", "File Number", "Ministry Area", "CCC Number", "Phone Number",
        "Sex", "Age", "Intake Date", "Programme Status", "Referral Source",
        "Caregiver", "Weight (kg)", "Height (cm)", "Diagnosis", "Household Children",
        "Orphans Supported", "Guardian", "Discharge Date", "Discharge Reason",
    ]
    widths = [26, 13, 16, 15, 14, 6, 6, 13, 18, 26, 24, 9, 9, 26, 10, 10, 16, 13, 22]
    title_block(ws, "Patient Register",
                "One row per patient. Do not merge cells. Delete the DEMO rows before importing live data.", len(headers))
    header_row(ws, 5, headers, widths)

    demo = [
        ["DEMO Patient 001", "LOM-001", "Nyalenda", "CCC/10001", "0700000001", "F", 34,
         "2026-03-14", "Active", "Demo County Referral Hospital", "Demo Caregiver", 58, 162,
         "HIV — on ART", 2, 1, "Single parent", "", ""],
        ["DEMO Patient 002", "LOM-002", "Ombeyi", "CCC/10002", "0700000002", "M", 47,
         "2026-01-08", "Discharged", "Demo CCC — Nyalenda Health Centre", "Demo Caregiver", 66, 170,
         "TB — treatment completed", 3, 2, "Grandparent", "2026-06-30", "Treatment completed; household stable"],
        ["DEMO Patient 003", "LOM-003", "Makongeni", "CCC/10003", "0700000003", "F", 29,
         "2026-05-02", "Awaiting verification", "Demo Sub-County Hospital (Homa Bay)", "Demo Caregiver", 51, 158,
         "HIV — newly initiated", 1, 0, "Both parents", "", ""],
    ]
    for i, row in enumerate(demo, start=6):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=v)
    style_data(ws, 6, 6 + len(demo) - 1, len(headers))
    for i in range(6, 9):
        ws.cell(row=i, column=1).font = Font(size=9, bold=True, color="9A5B00")

    dv_status = DataValidation(type="list", formula1='"Active,Enrolled,Awaiting verification,Discharged"', allow_blank=True)
    dv_sex = DataValidation(type="list", formula1='"F,M"', allow_blank=True)
    ws.add_data_validation(dv_status); ws.add_data_validation(dv_sex)
    dv_status.add(f"I6:I5000"); dv_sex.add("F6:F5000")

    # ---------------- Home Visits ----------------
    ws2 = wb.create_sheet("Home Visits")
    vh = ["File Number", "Patient Name", "Visit Date", "Visit Type", "Weight (kg)",
          "Blood Pressure", "Outcome", "Notes", "Next Visit Due"]
    title_block(ws2, "Home Visits",
                "One row per visit. Link to the patient using File Number where available.", len(vh))
    header_row(ws2, 5, vh, [13, 24, 13, 13, 10, 14, 16, 46, 13])
    visits = [
        ["LOM-001", "DEMO Patient 001", "2026-09-02", "Home visit", 59, "118/76", "Patient seen",
         "Adherence discussed; foodbag issued. Weight improved.", "2026-09-16"],
        ["LOM-001", "DEMO Patient 001", "2026-09-17", "Home visit", 60, "120/78", "Patient seen",
         "Household assessment completed. Referred for refill.", "2026-10-01"],
        ["LOM-002", "DEMO Patient 002", "2026-06-20", "Home visit", 66, "126/82", "Patient seen",
         "Discharge evaluation discussed with household.", ""],
        ["LOM-003", "DEMO Patient 003", "2026-09-28", "Home visit", 51, "", "Not found at home",
         "Patient absent; neighbour confirmed household present. Retry scheduled.", "2026-10-05"],
    ]
    for i, row in enumerate(visits, start=6):
        for j, v in enumerate(row, start=1):
            ws2.cell(row=i, column=j, value=v)
    style_data(ws2, 6, 6 + len(visits) - 1, len(vh))
    dv_out = DataValidation(type="list", formula1='"Patient seen,Not found at home,Declined visit,Referred for clinical review"', allow_blank=True)
    ws2.add_data_validation(dv_out); dv_out.add("G6:G5000")

    # ---------------- Hospital Referrals ----------------
    ws3 = wb.create_sheet("Hospital Referrals")
    hh = ["File Number", "Patient Name", "Request Date", "Facility", "Reason for Referral",
          "Priority", "Transport Required", "Estimated Fare (KES)", "Status", "Notes"]
    title_block(ws3, "Hospital Referrals & Transport",
                "One row per referral. Financial authorisation and payment are recorded separately in the DIMS.", len(hh))
    header_row(ws3, 5, hh, [13, 24, 13, 28, 40, 10, 16, 18, 16, 30])
    refs = [
        ["LOM-001", "DEMO Patient 001", "2026-09-20", "Demo County Referral Hospital",
         "ART refill and clinical review", "Routine", "Yes", 1200, "Approved",
         "Two persons travelling; one escort."],
        ["LOM-003", "DEMO Patient 003", "2026-10-02", "Demo Sub-County Hospital (Homa Bay)",
         "Persistent fever — urgent review", "Urgent", "Yes", 900, "Submitted", ""],
    ]
    for i, row in enumerate(refs, start=6):
        for j, v in enumerate(row, start=1):
            ws3.cell(row=i, column=j, value=v)
    style_data(ws3, 6, 6 + len(refs) - 1, len(hh))
    dv_pri = DataValidation(type="list", formula1='"Routine,Urgent,Critical"', allow_blank=True)
    ws3.add_data_validation(dv_pri); dv_pri.add("F6:F5000")

    # ---------------- Field Dictionary ----------------
    ws4 = wb.create_sheet("Field Dictionary")
    dd = ["Column", "Required", "Definition", "Example"]
    title_block(ws4, "Field Dictionary",
                "Definitions used by the import engine. Ambiguous columns are reported, never guessed silently.", len(dd))
    header_row(ws4, 5, dd, [26, 11, 72, 24])
    definitions = [
        ["Patient Name", "Yes", "Full name as written on the LOMAMO patient assessment form.", "DEMO Patient 001"],
        ["File Number", "No", "The organisation's patient file / registration number. Used for duplicate matching.", "LOM-001"],
        ["Ministry Area", "Yes", "The programme area of the patient. Must match a configured ministry area.", "Nyalenda"],
        ["CCC Number", "No", "Comprehensive Care Clinic identifier. The strongest duplicate-matching key.", "CCC/10001"],
        ["Phone Number", "No", "Contact number for follow-up. Treated as personal data.", "0700000001"],
        ["Sex", "No", "F or M.", "F"],
        ["Age", "No", "Age in completed years at intake.", "34"],
        ["Intake Date", "Yes", "Date of assessment / enrolment. Required for Patient-Days calculations.", "2026-03-14"],
        ["Programme Status", "No", "Active, Enrolled, Awaiting verification or Discharged.", "Active"],
        ["Referral Source", "No", "Hospital, CCC, PSC or other facility that referred the patient.", "Demo County Referral Hospital"],
        ["Caregiver", "No", "The field officer responsible for the patient.", "Demo Caregiver"],
        ["Weight (kg)", "No", "Most recent recorded weight. Used for trend alerts.", "58"],
        ["Height (cm)", "No", "Height at assessment.", "162"],
        ["Diagnosis", "No", "Primary clinical concern. Clinical verification is required after import.", "HIV — on ART"],
        ["Household Children", "No", "Number of children in the patient's household.", "2"],
        ["Orphans Supported", "No", "Children meeting the approved orphan-prevention outcome definition.", "1"],
        ["Guardian", "No", "Guardian arrangement for the household.", "Single parent"],
        ["Discharge Date", "No", "Date the discharge was approved. Overrides a conflicting status.", "2026-06-30"],
        ["Discharge Reason", "No", "Documented reason for discharge.", "Treatment completed"],
    ]
    for i, row in enumerate(definitions, start=6):
        for j, v in enumerate(row, start=1):
            ws4.cell(row=i, column=j, value=v)
    style_data(ws4, 6, 6 + len(definitions) - 1, len(dd))

    wb.save(FILE)
    print(f"Created {FILE.relative_to(ROOT)} — {FILE.stat().st_size/1024:,.1f} KB")
    return FILE


if __name__ == "__main__":
    build()
