#!/usr/bin/env python3
"""verify_no_patient_data.py - prove that nothing sensitive is published by
this repository: no patient identity, and no organisational secret.

It works in passes:

  PASS 1 - identities.  Reads the master workbook (from a path you give it, or
           from Git history if the file was removed from the working tree) and
           extracts every patient identity: a name on a row that also carries a
           date in column I, plus every telephone number in the Waiting List.
           A name that is only two tokens of the demo name pools is still
           treated as an identity - the check is deliberately strict.

  PASS 2 - the sweep.   Scans every text file in the working tree and the
           contents of every ZIP archive for those identities and numbers, and
           reports the files that match.

  PASS 3 - secrets.     Scans the same ground again for what a patient-identity
           sweep cannot see: vehicle registration plates, chassis numbers,
           handset IMEIs, and telephone numbers outside the demo band. The
           October 2026 inventory workbooks published all of these; see
           v6/docs/23_INVENTORY_REDACTION.md.

  PASS 4 - the register. Reads v6/src/inventory_data.js and asserts, row by
           row, that it carries structure and nothing else: no monetary value,
           no device identifier, no person used as a location, no narrative note
           mistaken for an item, and the count templates still present.

  PASS 5 - conventions. Every telephone number and mailbox seeded for a person
           must sit in reserved space (the 0700 demo band, the .example TLD), so
           an invented record can never carry a real person's contact details.

  PASS 6 - the files.   No Administration workbook may be tracked at the
           repository root. The blank demo-only import template under
           v6/examples is the one spreadsheet this repository publishes.

It also checks Git history for the workbook path, so "removed from the working
tree" is never mistaken for "removed from the repository".

Usage:
    python3 v6/tools/verify_no_patient_data.py
    python3 v6/tools/verify_no_patient_data.py --workbook /path/to/wb.xlsx
    python3 v6/tools/verify_no_patient_data.py --json

Exit code 0 means clean; 1 means something matched and is printed.

Requires: openpyxl for PASS 1  (python3 -m pip install --user openpyxl).
Passes 3 to 6 need nothing but Python.
"""
from __future__ import annotations

import argparse
import datetime
import json
import os
import re
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
WORKBOOK_NAME = "20260101 LOMAMO registration patients 2026 (1).xlsx"
AREA_SHEETS = ["Manyatta ", "Nyahera", "Nyalenda", "Ombeyi", "Nairobi Road", "Ahero", "Chiga"]
NOT_NAMES = {"manyatta", "nyahera", "nyalenda", "ombeyi", "nairobi road", "ahero", "chiga",
             "homabay", "waiting list", "total", "new", "basic", "name patient"}
BINARY_EXT = {"xlsx", "xls", "pdf", "docx", "doc", "png", "jpg", "jpeg", "webp", "gif", "zip", "gz", "bundle"}

# --------------------------------------------------------------------------
# PASS 3 - organisational secrets.
#
# The patient sweep above was not enough. The October 2026 inventory workbooks
# published vehicle registration plates, chassis numbers, handset IMEIs, staff
# SIM numbers, asset values and which member of staff carries which device.
# None of that is patient data, so a patient-identity sweep walked straight
# past it. These patterns are the ones that should have caught it.
# --------------------------------------------------------------------------
SECRET_PATTERNS = {
    "vehicle registration plate": re.compile(r"\bK[A-Z]{2}\s?\d{3}\s?[A-Z]\b"),
    "handset IMEI": re.compile(r"\b3\d{14}\b"),
    "mobile number outside the demo band": re.compile(
        r"\b(?:0|\+?254)?7[1-9]\d{7}\b|\b01\d{8}\b"),
    "chassis or VIN": re.compile(r"chassis[^A-Z0-9]{0,14}[A-HJ-NPR-Z0-9]{11,}", re.I),
    "landline or subscriber number": re.compile(r"\b0\d{8,9}\b"),
}
# Invented on purpose, documented in the source, and impossible to mistake for
# a real identifier: the demo fleet in v6/src/data.js and the PRNG constants.
SECRET_ALLOWLIST = {
    "KDX 001A", "KDG 442B", "KCC 887C",        # demo vehicles, data.js FLEET seed
    "1013904223", "2654435761",                # linear-congruential constants
    # Fixtures in the acceptance suites. They are invented numbers used to prove
    # that MESSAGING.screen and the referral form REFUSE contact details; the
    # control only means something if the fixture has a real number's shape.
    "0712345678", "0112345678", "254712345678", "0722111222",
}
# Files that quote the patterns in order to detect them.
# Only the two tools that carry the patterns themselves are skipped. The
# documentation is held to the same standard as the code: a document that quotes
# a plate or an IMEI re-publishes it, so this verifier would flag its own
# runbook. That is deliberate.
SECRET_SKIP = {"v6/tools/verify_no_patient_data.py", "v6/tools/redact_inventory.py"}
# The demo estate owns these bands. Anything seeded for a person must sit in
# one of them, so a fabricated record can never carry a real identifier.
DEMO_PHONE = re.compile(r"^0700\d{6}$")
DEMO_EMAIL_DOMAINS = ("@lomamo.example", "@lomamo.demo", "@example.com", "@example.org")
INTERNAL_TAG = re.compile(r"^(?:LMM|LMMK)[A-Z0-9/\-]*$")


def sh(args: list[str], **kw) -> subprocess.CompletedProcess:
    return subprocess.run(args, cwd=ROOT, capture_output=True, text=True, **kw)


def workbook_from_history() -> Path | None:
    """Find the workbook in Git history and write it to a temporary file."""
    commits = sh(["git", "log", "--all", "--format=%H", "--", WORKBOOK_NAME]).stdout.split()
    if not commits:
        return None
    blob = sh(["git", "show", f"{commits[0]}:{WORKBOOK_NAME}"], check=False)
    if blob.returncode != 0:
        # binary content: use cat-file instead of show
        sha = sh(["git", "rev-parse", f"{commits[0]}:{WORKBOOK_NAME}"]).stdout.strip()
        raw = subprocess.run(["git", "cat-file", "blob", sha], cwd=ROOT, capture_output=True)
        if raw.returncode != 0:
            return None
        tmp = Path(tempfile.gettempdir()) / "cdims-workbook-from-history.xlsx"
        tmp.write_bytes(raw.stdout)
        return tmp
    return None


def extract_identities(path: Path) -> tuple[set[str], set[str]]:
    try:
        import openpyxl  # noqa: WPS433 - optional dependency, imported late
    except ImportError:
        print("openpyxl is not installed: python3 -m pip install --user openpyxl")
        print("Running the history and archive checks only.")
        return set(), set()
    if not path.exists() or path.stat().st_size < 1024:
        print(f"the workbook is not readable at {path} (missing or truncated); identity sweep skipped")
        return set(), set()
    import warnings
    warnings.filterwarnings("ignore")
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    except Exception as exc:
        print(f"the workbook could not be opened ({type(exc).__name__}); identity sweep skipped")
        return set(), set()
    names: set[str] = set()
    phones: set[str] = set()
    for sheet in AREA_SHEETS:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for row in ws.iter_rows(min_row=4, max_col=9, values_only=True):
            name, date_in = row[0], row[8]
            if isinstance(name, str) and isinstance(date_in, (datetime.datetime, datetime.date)):
                n = " ".join(name.split())
                if n.lower() not in NOT_NAMES and re.fullmatch(r"[A-Za-z'.\- ]{6,60}", n) and len(n.split()) >= 2:
                    names.add(n)
    if "Waiting List" in wb.sheetnames:
        ws = wb["Waiting List"]
        for row in ws.iter_rows(min_row=3, max_col=2, values_only=True):
            name, phone = row[0], row[1]
            if isinstance(name, str) and len(name.split()) >= 2 and name.lower().strip() not in NOT_NAMES:
                names.add(" ".join(name.split()))
            if phone is not None:
                digits = re.sub(r"\D", "", str(phone))
                if 8 <= len(digits) <= 10:
                    phones.add(digits)
    wb.close()
    return names, phones


def staff_names() -> set[str]:
    """Names that legitimately appear in the repository: the official roster."""
    out = set()
    src = ROOT / "v6" / "src"
    for f in ("data6.js", "org_teams.js", "portal_content.js"):
        p = src / f
        if p.exists():
            out |= {m.lower() for m in re.findall(r"name:\s*'([^']+)'", p.read_text())}
            out |= {m.lower() for m in re.findall(r"\{ name:'([^']+)'", p.read_text())}
    return out


def iter_text_files():
    for dirpath, dirnames, filenames in os.walk(ROOT):
        dirnames[:] = [d for d in dirnames if d not in {".git", "node_modules", "__pycache__"}]
        for fn in filenames:
            ext = fn.lower().rsplit(".", 1)[-1] if "." in fn else ""
            if ext in BINARY_EXT:
                continue
            yield Path(dirpath) / fn


def scan_zip(path: Path, name_rx, phone_rx):
    hits = {"names": 0, "phones": 0, "files": 0, "inner_xlsx": []}
    try:
        with zipfile.ZipFile(path) as z:
            for info in z.infolist():
                if info.is_dir():
                    continue
                inner = Path(info.filename)
                if inner.suffix.lower() in {".xlsx", ".xls"}:
                    hits["inner_xlsx"].append(info.filename)
                    continue
                if inner.suffix.lower() in BINARY_EXT:
                    continue
                try:
                    txt = z.read(info).decode("utf-8", errors="ignore")
                except Exception:
                    continue
                hits["files"] += 1
                hits["names"] += len(name_rx.findall(txt))
                hits["phones"] += len(phone_rx.findall(txt))
    except Exception as exc:  # unreadable archive is reported, not skipped silently
        hits["error"] = str(exc)[:120]
    return hits


# --------------------------------------------------------------------------
# PASS 3 - the sweep for organisational secrets
# --------------------------------------------------------------------------
STAFF_TOKENS = {
    "washington", "ashar", "catherine", "denis", "daniel", "julia", "julie",
    "martin", "musa", "moses", "oscar", "evans", "wickliffe", "wicklyff",
    "john", "irene", "iscah", "millicent", "caroline", "maureen", "chamah",
    "otieno", "odhiambo", "omondi", "obila", "lubalo", "wayodi", "oguta",
    "omolo", "odero", "onyango", "achieng", "abuto", "nyanjong", "atieno",
    "awuor", "anyango", "shikuku", "mildred", "naphtaly", "oduja", "mito",
    "mcclanahan",
}


def _hits(txt: str) -> dict:
    out = {}
    for label, rx in SECRET_PATTERNS.items():
        found = [m.group(0).strip() for m in rx.finditer(txt)]
        found = sorted({h for h in found
                        if h not in SECRET_ALLOWLIST and not DEMO_PHONE.match(h)})
        if found:
            out[label] = found[:4]
    return out


def scan_secrets() -> tuple[list[dict], int]:
    """Working tree and the contents of every archive, for secrets."""
    offenders, scanned = [], 0
    for f in iter_text_files():
        rel = str(f.relative_to(ROOT))
        if rel in SECRET_SKIP:
            continue
        try:
            txt = f.read_bytes().decode("utf-8", errors="ignore")
        except Exception:
            continue
        scanned += 1
        h = _hits(txt)
        if h:
            offenders.append({"file": rel, **h})
    for f in ROOT.rglob("*.zip"):
        if ".git" in f.parts:
            continue
        blob = ""
        try:
            with zipfile.ZipFile(f) as z:
                for info in z.infolist():
                    if info.is_dir() or Path(info.filename).suffix.lower() in BINARY_EXT:
                        continue
                    try:
                        blob += z.read(info).decode("utf-8", errors="ignore")
                    except Exception:
                        continue
        except Exception as exc:
            offenders.append({"file": str(f.relative_to(ROOT)), "unreadable archive": str(exc)[:80]})
            continue
        h = _hits(blob)
        if h:
            offenders.append({"file": str(f.relative_to(ROOT)) + " (archive)", **h})
    return offenders, scanned


def _field(row: str, name: str):
    m = re.search(name + r':(?:"((?:[^"\\]|\\.)*)"|(-?\d+))', row)
    if not m:
        return None
    return m.group(1) if m.group(1) is not None else int(m.group(2))


def inventory_structure() -> list[str]:
    """PASS 4 - the published register carries structure and nothing else."""
    fails: list[str] = []
    p = ROOT / "v6" / "src" / "inventory_data.js"
    if not p.exists():
        return ["v6/src/inventory_data.js is missing"]
    txt = p.read_text()
    m = re.search(r"const INVENTORY_SEED = \[(.*?)\n\];", txt, re.S)
    if not m:
        return ["INVENTORY_SEED not found in v6/src/inventory_data.js"]
    rows = re.findall(r"\{(.*?)\}", m.group(1), re.S)
    if len(rows) < 100:
        fails.append(f"only {len(rows)} register rows parsed - expected about 122")
    for i, r in enumerate(rows, 1):
        value = _field(r, "value")
        serial = _field(r, "serial")
        loc = _field(r, "loc") or ""
        item = _field(r, "item") or ""
        qty = _field(r, "qty")
        if value not in ("", None):
            fails.append(f"row {i} ({item}) carries a monetary value: {value!r}")
        if serial and not INTERNAL_TAG.match(serial):
            fails.append(f"row {i} ({item}) carries a device identifier: {serial!r}")
        if re.fullmatch(r"\d+", item):
            fails.append(f"row {i} is a narrative note, not an item")
        if isinstance(qty, int) and qty >= 1000:
            fails.append(f"row {i} ({item}) has an implausible quantity {qty} - a note row was read as an item")
        bad = [t for t in re.findall(r"[A-Za-z']+", loc) if t.lower() in STAFF_TOKENS]
        if bad:
            fails.append(f"row {i} locates an asset at a person: {loc!r}")
    if "function inventorySeed(" not in txt:
        fails.append("inventorySeed() is missing - the register would render empty")
    for const in ("INVENTORY_CATEGORIES", "INVENTORY_SITES", "INVENTORY_REDACTION",
                  "INVENTORY_NOTES", "KITCHEN_CATALOG", "KITCHEN_COUNT_COLUMNS",
                  "SUPPLY_CATALOG", "SUPPLY_COUNT_COLUMNS", "INVENTORY_SOURCE"):
        if f"const {const}" not in txt:
            fails.append(f"{const} is missing - the workbook's structure was not carried over")
    return fails


EMAIL_RX = re.compile(r"\b[\w.+-]+@[\w-]+\.[A-Za-z]{2,}\b")
PHONE_LITERAL = re.compile(r"phone:\s*'([^']*)'")
EMAIL_PUBLIC_OK = {"contactus@lovemattersmost.org"}


def contact_conventions() -> list[str]:
    """PASS 5 - every seeded person is reachable only in reserved space."""
    fails: list[str] = []
    for f in sorted((ROOT / "v6" / "src").glob("*.js")):
        txt = f.read_text()
        rel = f"v6/src/{f.name}"
        for n in PHONE_LITERAL.findall(txt):
            if n and not DEMO_PHONE.match(n) and not n == "0700":
                fails.append(f"{rel} seeds a telephone number outside the demo band: {n!r}")
        for gen in re.findall(r"'07'\s*\+", txt):
            fails.append(f"{rel} still builds a number from '07' + a random range")
        for e in EMAIL_RX.findall(txt):
            if e in EMAIL_PUBLIC_OK:
                continue
            if not e.endswith(DEMO_EMAIL_DOMAINS):
                fails.append(f"{rel} seeds a mailbox on a live domain: {e!r}")
    return fails


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", help="path to the master workbook (default: find it in Git history)")
    ap.add_argument("--json", action="store_true", help="machine-readable output")
    args = ap.parse_args()

    report = {"repository": str(ROOT), "checks": [], "clean": True}

    def add(name, ok, detail=""):
        report["checks"].append({"check": name, "ok": bool(ok), "detail": detail})
        if not ok:
            report["clean"] = False

    # --- 1. history check ---------------------------------------------------
    hist = sh(["git", "log", "--all", "--oneline", "--", WORKBOOK_NAME]).stdout.strip()
    tree = sh(["git", "ls-tree", "-r", "HEAD", "--name-only"]).stdout
    add("workbook absent from the working tree of HEAD", WORKBOOK_NAME not in tree,
        "still tracked in HEAD" if WORKBOOK_NAME in tree else "")
    add("workbook absent from all Git history", not hist,
        f"reachable in {len(hist.splitlines())} commit(s) - run v6/tools/purge_patient_data.sh" if hist else "")
    depth = sh(["git", "rev-list", "--all", "--count"]).stdout.strip()
    shallow = (ROOT / ".git" / "shallow").exists()
    if shallow or depth == "1":
        report["checks"].append({"check": "NOTE shallow clone", "ok": True,
                                 "detail": f"only {depth} commit(s) reachable locally; history findings are inconclusive - unshallow before trusting them"})

    # --- 2. identities ------------------------------------------------------
    wb_path = Path(args.workbook) if args.workbook else None
    if wb_path is None:
        wb_path = ROOT / WORKBOOK_NAME if (ROOT / WORKBOOK_NAME).exists() else workbook_from_history()
    names, phones = (set(), set())
    if wb_path and wb_path.exists():
        names, phones = extract_identities(wb_path)
        names = {n for n in names if n.lower() not in staff_names()}
        add("identities extracted from the workbook", bool(names or phones),
            f"{len(names)} patient names, {len(phones)} telephone numbers")
    else:
        report["checks"].append({"check": "NOTE workbook not available", "ok": True,
                                 "detail": "no local copy and none in this clone's history; the identity sweep was skipped"})

    # --- 3. the sweep -------------------------------------------------------
    if names or phones:
        name_rx = re.compile(r"\b(" + "|".join(re.escape(n) for n in sorted(names)) + r")\b", re.I)
        phone_rx = re.compile(r"(?:\+?254|0)?(" + "|".join(re.escape(p[-9:]) for p in sorted(phones)) + r")\b")
        offenders, scanned = [], 0
        for f in iter_text_files():
            try:
                txt = f.read_bytes().decode("utf-8", errors="ignore")
            except Exception:
                continue
            scanned += 1
            n = len(name_rx.findall(txt))
            p = len(phone_rx.findall(txt))
            if n or p:
                offenders.append({"file": str(f.relative_to(ROOT)), "names": n, "phones": p})
        add("no patient identity in any working-tree file", not offenders,
            f"{scanned} files scanned; matches: " + ", ".join(o["file"] for o in offenders[:6]))
        zip_offenders = []
        for f in ROOT.rglob("*.zip"):
            if ".git" in f.parts:
                continue
            h = scan_zip(f, name_rx, phone_rx)
            if h.get("names") or h.get("phones") or h.get("inner_xlsx") or h.get("error"):
                zip_offenders.append({"file": str(f.relative_to(ROOT)), **h})
        add("no patient identity inside any archive", not zip_offenders, json.dumps(zip_offenders[:4]))
        report["scanned_files"] = scanned
        report["identities"] = {"names": len(names), "phones": len(phones)}

    # --- 4. organisational secrets ------------------------------------------
    offenders, secret_scanned = scan_secrets()
    add("no vehicle plate, IMEI, chassis number or live-range telephone number "
        "in any file or archive", not offenders,
        f"{secret_scanned} files scanned; " +
        "; ".join(f"{o['file']} ({', '.join(k for k in o if k != 'file')})"
                  for o in offenders[:5]))
    report["secret_scan"] = offenders[:10]

    # --- 5. the published register -------------------------------------------
    inv_fails = inventory_structure()
    add("the published equipment register carries structure only", not inv_fails,
        "; ".join(inv_fails[:5]) if inv_fails else
        "no values, no device identifiers, no person as a location, no note rows")

    # --- 6. contact conventions ----------------------------------------------
    contact_fails = contact_conventions()
    add("every seeded person is reachable only in reserved space", not contact_fails,
        "; ".join(contact_fails[:5]) if contact_fails else
        "telephone numbers in the 0700 demo band, mailboxes on .example")

    # --- 7. the workbooks themselves -----------------------------------------
    tracked = sh(["git", "ls-files"]).stdout.splitlines()
    rooted = [t for t in tracked
              if t.lower().endswith((".xlsx", ".xls")) and "/" not in t]
    add("no Administration workbook is tracked at the repository root", not rooted,
        ", ".join(rooted) if rooted else
        "structure only; the files stay on the office host")

    if args.json:
        print(json.dumps(report, indent=1))
    else:
        print("LOMAMO C-DIMS · personal-data verification")
        print("-" * 64)
        for c in report["checks"]:
            mark = "ok  " if c["ok"] else "FAIL"
            print(f"  [{mark}] {c['check']}" + (f" — {c['detail']}" if c["detail"] else ""))
        print("-" * 64)
        print("CLEAN" if report["clean"] else "NOT CLEAN — see the failures above")
    return 0 if report["clean"] else 1


if __name__ == "__main__":
    sys.exit(main())
